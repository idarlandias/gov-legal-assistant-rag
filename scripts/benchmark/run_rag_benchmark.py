"""Roda as perguntas do benchmark da Fase 3 no RAG atual e registra os resultados.

Salva sempre um JSON em docs/fase3/benchmark/results/. Com --xlsx, preenche também as
linhas "RAG atual" da planilha de coleta. As notas da rubrica ficam para avaliação humana.

Uso:
    python scripts/benchmark/run_rag_benchmark.py --dry-run
    python scripts/benchmark/run_rag_benchmark.py --ids 1,2,11 --xlsx
    python scripts/benchmark/run_rag_benchmark.py --xlsx --overwrite
"""

from __future__ import annotations

import argparse
import json
import os
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

BENCH_DIR = ROOT / "docs" / "fase3" / "benchmark"
PERGUNTAS_JSON = BENCH_DIR / "perguntas.json"
XLSX_PATH = BENCH_DIR / "planilha_benchmark.xlsx"
RESULTS_DIR = BENCH_DIR / "results"
RAG_TOOL = "RAG atual"


def load_perguntas(ids: set[int] | None) -> list[dict]:
    data = json.loads(PERGUNTAS_JSON.read_text(encoding="utf-8"))
    perguntas = data["perguntas"]
    if ids:
        faltando = ids - {p["id"] for p in perguntas}
        if faltando:
            raise SystemExit(f"ERRO: IDs inexistentes em perguntas.json: {sorted(faltando)}")
        perguntas = [p for p in perguntas if p["id"] in ids]
    return perguntas


def git_commit() -> str | None:
    try:
        return subprocess.run(["git", "rev-parse", "--short", "HEAD"], cwd=ROOT,
                              capture_output=True, text=True, check=True).stdout.strip()
    except (OSError, subprocess.CalledProcessError):
        return None


def instrument_usage(pipeline: Any) -> dict[str, int]:
    """Envolve _call_chat_completions para somar tokens de todas as chamadas (inclui tools)."""
    usage = {"prompt_tokens": 0, "completion_tokens": 0, "calls": 0, "calls_sem_usage": 0}
    original = pipeline._call_chat_completions

    def wrapped(**kwargs: Any) -> Any:
        resp = original(**kwargs)
        usage["calls"] += 1
        u = getattr(resp, "usage", None)
        if u is None:
            usage["calls_sem_usage"] += 1
        else:
            usage["prompt_tokens"] += getattr(u, "prompt_tokens", 0) or 0
            usage["completion_tokens"] += getattr(u, "completion_tokens", 0) or 0
        return resp

    pipeline._call_chat_completions = wrapped
    return usage


def fmt_sources(sources: list) -> str:
    vistos = dict.fromkeys(f"{s}:p{p}" for s, p in sources)
    return "; ".join(vistos)


def run(perguntas: list[dict], k: int, domain: str, sleep_s: float) -> tuple[dict, list[dict]]:
    from dotenv import load_dotenv

    load_dotenv(ROOT / ".env")
    from src.observability.trace import trace
    from src.pipeline.rag import RAGPipeline

    # RAGPipeline direto (não build_rag_pipeline) para nunca resetar/reindexar o banco aqui
    pipeline = RAGPipeline(corpus_dir=str(ROOT / "data" / "corpus"),
                           persist_dir=str(ROOT / "data" / "chroma"))
    n_chunks = pipeline.collection.count()
    if n_chunks == 0:
        raise SystemExit("ERRO: coleção Chroma vazia. Indexe o corpus antes (rode o app uma vez).")

    usage = instrument_usage(pipeline)
    meta = {
        "ferramenta": RAG_TOOL,
        "data": datetime.now().isoformat(timespec="seconds"),
        "provider": os.environ.get("LLM_PROVIDER", "gemini").lower(),
        "llm_model": pipeline.llm_model,
        "embed_model": pipeline.embed_model,
        "k": k,
        "domain": domain,
        "chunks_indexados": n_chunks,
        "git_commit": git_commit(),
    }

    resultados = []
    for i, p in enumerate(perguntas):
        before = dict(usage)
        rec: dict[str, Any] = {"id": p["id"], "categoria": p["categoria"], "pergunta": p["pergunta"]}
        t0 = time.perf_counter()
        try:
            with trace("benchmark_fase3", pergunta_id=p["id"]):
                out = pipeline.answer(p["pergunta"], k=k, domain=domain)
            rec["resposta"] = out.get("answer", "")
            rec["fontes"] = fmt_sources(out.get("sources", []))
            rec["erro"] = None
        except Exception as e:  # registra e segue; o exit code sinaliza a falha
            rec["resposta"], rec["fontes"], rec["erro"] = "", "", f"{type(e).__name__}: {e}"
        rec["latencia_s"] = round(time.perf_counter() - t0, 2)
        tem_usage = usage["calls"] > before["calls"] and usage["calls_sem_usage"] == before["calls_sem_usage"]
        rec["tokens_entrada"] = usage["prompt_tokens"] - before["prompt_tokens"] if tem_usage else None
        rec["tokens_saida"] = usage["completion_tokens"] - before["completion_tokens"] if tem_usage else None
        resultados.append(rec)

        status = "ERRO " + rec["erro"] if rec["erro"] else f"ok {rec['latencia_s']}s"
        print(f"[{i + 1}/{len(perguntas)}] #{p['id']} {status}", file=sys.stderr)
        if sleep_s and i < len(perguntas) - 1:
            time.sleep(sleep_s)
    return meta, resultados


def fill_xlsx(path: Path, meta: dict, resultados: list[dict], overwrite: bool) -> tuple[int, int]:
    from openpyxl import load_workbook

    wb = load_workbook(path)  # sem data_only: preserva as fórmulas
    ws = wb["Respostas"]
    linhas = {(row[0].value, row[1].value): row[0].row
              for row in ws.iter_rows(min_row=2, max_col=2) if row[0].value is not None}
    versao = f"{meta['provider']}/{meta['llm_model']} k={meta['k']} ({meta['git_commit'] or 's/ commit'})"
    escritas = puladas = 0
    for rec in resultados:
        if rec["erro"]:
            continue
        r = linhas.get((rec["id"], RAG_TOOL))
        if r is None:
            print(f"AVISO: pergunta #{rec['id']} não existe na planilha; rode build_planilha.py",
                  file=sys.stderr)
            puladas += 1
            continue
        if ws[f"G{r}"].value and not overwrite:
            print(f"AVISO: #{rec['id']} já tem resposta na planilha; use --overwrite", file=sys.stderr)
            puladas += 1
            continue
        ws[f"E{r}"] = meta["data"][:10]
        ws[f"F{r}"] = versao
        ws[f"G{r}"] = rec["resposta"]
        ws[f"H{r}"] = rec["fontes"]
        ws[f"I{r}"] = rec["latencia_s"]
        ws[f"J{r}"] = rec["tokens_entrada"]
        ws[f"K{r}"] = rec["tokens_saida"]
        escritas += 1
    try:
        wb.save(path)
    except PermissionError:
        raise SystemExit(f"ERRO: não consegui salvar {path}. Feche a planilha no Excel e rode de novo "
                         "(o JSON de resultados já foi salvo).")
    return escritas, puladas


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--ids", help="IDs separados por vírgula (padrão: todas)")
    ap.add_argument("--k", type=int, default=5, help="chunks recuperados (padrão 5)")
    ap.add_argument("--domain", default="auto", help="filtro de domínio do Chroma (padrão auto)")
    ap.add_argument("--sleep", type=float, default=2.0, help="pausa entre perguntas, em s")
    ap.add_argument("--xlsx", action="store_true", help="preenche as linhas 'RAG atual' da planilha")
    ap.add_argument("--overwrite", action="store_true", help="sobrescreve respostas já preenchidas")
    ap.add_argument("--dry-run", action="store_true", help="só lista as perguntas, sem chamar o LLM")
    args = ap.parse_args()

    ids = {int(x) for x in args.ids.split(",")} if args.ids else None
    perguntas = load_perguntas(ids)

    if args.dry_run:
        for p in perguntas:
            print(f"#{p['id']:>2} [{p['categoria']}] {p['pergunta']}")
        return 0
    if args.xlsx and not XLSX_PATH.exists():
        raise SystemExit(f"ERRO: {XLSX_PATH} não existe. Rode scripts/benchmark/build_planilha.py")

    meta, resultados = run(perguntas, args.k, args.domain, args.sleep)

    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    out = RESULTS_DIR / f"rag-{datetime.now():%Y%m%d-%H%M%S}.json"
    out.write_text(json.dumps({"meta": meta, "resultados": resultados}, ensure_ascii=False, indent=2),
                   encoding="utf-8")
    print(f"Resultados: {out}", file=sys.stderr)

    if args.xlsx:
        escritas, puladas = fill_xlsx(XLSX_PATH, meta, resultados, args.overwrite)
        print(f"Planilha: {escritas} linhas escritas, {puladas} puladas", file=sys.stderr)

    erros = [r for r in resultados if r["erro"]]
    sem_tokens = [r["id"] for r in resultados if not r["erro"] and r["tokens_entrada"] is None]
    if sem_tokens:
        print(f"AVISO: provider não retornou usage para {sem_tokens}; custo ficará em branco",
              file=sys.stderr)
    if erros:
        print(f"FALHAS: {len(erros)}/{len(resultados)} perguntas: {[r['id'] for r in erros]}",
              file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
