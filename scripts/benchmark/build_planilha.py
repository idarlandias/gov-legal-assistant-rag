"""Gera a planilha de coleta do benchmark da Fase 3 a partir de perguntas.json.

Uso:
    python scripts/benchmark/build_planilha.py            # cria se nao existir
    python scripts/benchmark/build_planilha.py --force    # sobrescreve (perde o preenchido!)
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
BENCH_DIR = ROOT / "docs" / "fase3" / "benchmark"
PERGUNTAS_JSON = BENCH_DIR / "perguntas.json"
XLSX_PATH = BENCH_DIR / "planilha_benchmark.xlsx"

FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True)
F_INPUT = Font(name=FONT, size=10, color="0000FF")
FILL_HEAD = PatternFill("solid", fgColor="1F3864")
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")
FILL_FORMULA = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

RAG_TOOL = "RAG atual"


def is_rag(tool: str) -> bool:
    """Ferramentas nossas ("RAG atual", "RAG Fase 3"): custo vem dos tokens, não de assinatura."""
    return tool.startswith("RAG")

# Colunas da aba Respostas (letra, cabecalho, largura, tipo)
RESP_COLS = [
    ("A", "ID", 6, "key"),
    ("B", "Ferramenta", 14, "key"),
    ("C", "Categoria", 16, "formula"),
    ("D", "Pergunta", 50, "formula"),
    ("E", "Data da coleta", 13, "input"),
    ("F", "Versão / plano", 18, "input"),
    ("G", "Resposta integral", 70, "input"),
    ("H", "Fontes citadas", 30, "input"),
    ("I", "Latência (s)", 11, "input"),
    ("J", "Tokens entrada", 11, "input"),
    ("K", "Tokens saída", 11, "input"),
    ("L", "Custo estimado (R$)", 13, "formula"),
    ("M", "Veredito (0-2)", 10, "score"),
    ("N", "Fundamentação (0-2)", 13, "score"),
    ("O", "Norma CE (0-2)", 10, "score"),
    ("P", "Acionável (0-2)", 10, "score"),
    ("Q", "Sem alucinação (0-2)", 13, "score"),
    ("R", "Nota (0-10)", 10, "formula"),
    ("S", "Observações", 40, "input"),
]

RUBRICA = [
    ("Veredito", "errado", "ambíguo", "correto"),
    ("Fundamentação", "sem citação", "lei genérica", "dispositivo específico correto"),
    ("Norma CE", "ignora", "menciona", "aplica corretamente"),
    ("Acionável", "só teoria", "orientação", "ato/minuta utilizável"),
    ("Sem alucinação", "inventa norma", "impreciso", "nenhuma"),
]


def _header(ws, row: int, labels: list[str]) -> None:
    for col, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=col, value=label)
        c.font, c.fill, c.border = F_HEAD, FILL_HEAD, BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


def _style(c, kind: str) -> None:
    c.border = BORDER
    c.alignment = WRAP
    if kind in ("input", "score"):
        c.font, c.fill = F_INPUT, FILL_INPUT
    elif kind == "formula":
        c.font, c.fill = F_BASE, FILL_FORMULA
    else:
        c.font = F_BASE


def build_leia_me(wb: Workbook, n_perg: int, ferramentas: list[str]) -> None:
    ws = wb.active
    ws.title = "LEIA-ME"
    ws.column_dimensions["A"].width = 24
    for col in "BCDE":
        ws.column_dimensions[col].width = 30

    ws["A1"] = "Benchmark Fase 3: RAG jurídico × concorrentes"
    ws["A1"].font = F_TITLE
    linhas = [
        f"Objetivo: medir se {', '.join(t for t in ferramentas if not is_rag(t))} já respondem bem às {n_perg} "
        "perguntas de balcão (Registro de Imóveis × LAI × LGPD). O resultado alimenta o Gate 1.",
        "1. Aba Perguntas: o oficial/DPO preenche o gabarito (colunas amarelas) ANTES de avaliar.",
        "2. Aba Respostas: uma linha por pergunta × ferramenta. Cole a resposta integral, sem editar.",
        "3. Mesma redação, sessão nova, sem contexto extra em cada ferramenta.",
        "4. As linhas 'RAG atual' e 'RAG Fase 3' são preenchidas pelo script run_rag_benchmark.py (--xlsx).",
        "5. Dê a nota de 0 a 2 em cada critério comparando com o gabarito. A Nota (0-10) só "
        "aparece com os 5 critérios preenchidos.",
        "6. Aba Parametros: preencha a assinatura mensal dos concorrentes e confira preço/câmbio.",
        "7. A aba Resumo e o sinal do Gate 1 são calculados sozinhos.",
    ]
    for i, txt in enumerate(linhas, start=3):
        ws.cell(row=i, column=1, value=txt).font = F_BASE

    r = 3 + len(linhas) + 1
    ws.cell(row=r, column=1, value="Legenda").font = F_BOLD
    c = ws.cell(row=r + 1, column=1, value="Preencher")
    c.fill, c.font = FILL_INPUT, F_INPUT
    ws.cell(row=r + 1, column=2, value="Célula de entrada (texto azul, fundo amarelo)").font = F_BASE
    c = ws.cell(row=r + 2, column=1, value="Fórmula")
    c.fill, c.font = FILL_FORMULA, F_BASE
    ws.cell(row=r + 2, column=2, value="Calculado; não editar").font = F_BASE

    r += 4
    ws.cell(row=r, column=1, value="Rubrica (0 a 2 por critério)").font = F_BOLD
    _header(ws, r + 1, ["Critério", "0", "1", "2"])
    for i, row in enumerate(RUBRICA, start=r + 2):
        for j, v in enumerate(row, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font, c.border, c.alignment = F_BASE, BORDER, WRAP

    r += len(RUBRICA) + 3
    ws.cell(row=r, column=1, value="Exemplo de linha preenchida (formato esperado)").font = F_BOLD
    exemplo = [
        ("ID / Ferramenta", "6 / Jus IA"),
        ("Data da coleta", "2026-09-25"),
        ("Versão / plano", "Plano Profissional (set/2026)"),
        ("Resposta integral", "Não. A informação sobre titularidade é fornecida por certidão..."),
        ("Fontes citadas", "Lei 6.015/73, art. 16-17; LGPD art. 7º"),
        ("Latência (s)", "9"),
        ("Notas 0-2", "Veredito 2 · Fundamentação 1 · Norma CE 0 · Acionável 1 · Alucinação 2 → Nota 6"),
        ("Observações", "Não citou o Código de Normas do CE nem gerou minuta."),
    ]
    for i, (k, v) in enumerate(exemplo, start=r + 1):
        ws.cell(row=i, column=1, value=k).font = F_BOLD
        ws.cell(row=i, column=2, value=v).font = F_BASE


def build_perguntas(wb: Workbook, perguntas: list[dict]) -> None:
    ws = wb.create_sheet("Perguntas")
    cols = [
        ("ID", 6, "key"), ("Categoria", 16, "key"), ("Pergunta", 70, "key"),
        ("Veredito esperado", 30, "input"), ("Base legal do gabarito", 40, "input"),
        ("Validado por", 20, "input"), ("Data da validação", 14, "input"),
    ]
    _header(ws, 1, [c[0] for c in cols])
    for idx, (_, w, _) in enumerate(cols, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = w
    for r, p in enumerate(perguntas, start=2):
        vals = [p["id"], p["categoria"], p["pergunta"], None, None, None, None]
        for cidx, (v, (_, _, kind)) in enumerate(zip(vals, cols), start=1):
            c = ws.cell(row=r, column=cidx, value=v)
            _style(c, kind)
    ws.freeze_panes = "D2"


def build_respostas(wb: Workbook, perguntas: list[dict], ferramentas: list[str]) -> int:
    ws = wb.create_sheet("Respostas")
    _header(ws, 1, [c[1] for c in RESP_COLS])
    for letter, _, width, _ in RESP_COLS:
        ws.column_dimensions[letter].width = width
    ws.row_dimensions[1].height = 30

    last_q = len(perguntas) + 1
    row = 2
    for tool in ferramentas:
        for p in perguntas:
            for letter, _, _, kind in RESP_COLS:
                c = ws[f"{letter}{row}"]
                if letter == "A":
                    c.value = p["id"]
                elif letter == "B":
                    c.value = tool
                elif letter == "C":
                    c.value = (f"=INDEX(Perguntas!$B$2:$B${last_q},"
                               f"MATCH(A{row},Perguntas!$A$2:$A${last_q},0))")
                elif letter == "D":
                    c.value = (f"=INDEX(Perguntas!$C$2:$C${last_q},"
                               f"MATCH(A{row},Perguntas!$A$2:$A${last_q},0))")
                elif letter == "L":
                    c.value = (f'=IF(COUNT(J{row}:K{row})=2,'
                               f"(J{row}*Parametros!$B$3+K{row}*Parametros!$B$4)"
                               f'/1000000*Parametros!$B$5,"")')
                    c.number_format = "0.0000"
                elif letter == "R":
                    c.value = f'=IF(COUNT(M{row}:Q{row})=5,SUM(M{row}:Q{row}),"")'
                _style(c, kind)
            row += 1
    last = row - 1

    dv_score = DataValidation(type="list", formula1='"0,1,2"', allow_blank=True,
                              showErrorMessage=True, error="Use 0, 1 ou 2")
    ws.add_data_validation(dv_score)
    dv_score.add(f"M2:Q{last}")

    ws.conditional_formatting.add(
        f"R2:R{last}",
        CellIsRule(operator="greaterThanOrEqual", formula=["8"],
                   fill=PatternFill("solid", fgColor="C6EFCE")))
    ws.conditional_formatting.add(
        f"R2:R{last}",
        CellIsRule(operator="lessThan", formula=["6"],
                   fill=PatternFill("solid", fgColor="FFC7CE")))
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:S{last}"
    return last


def build_parametros(wb: Workbook, ferramentas: list[str]) -> None:
    ws = wb.create_sheet("Parametros")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 60
    _header(ws, 1, ["Parâmetro", "Valor", "", "Fonte / observação"])
    params = [
        ("Modelo do RAG", "Groq openai/gpt-oss-20b / 120b",
         "CHEAP/PREMIUM_MODEL do .env; o script registra o modelo real em 'Versão / plano'"),
        ("Preço entrada (US$ / 1M tokens)", 0.15,
         "Premissa: preço do gpt-oss-120b na Groq (teto; o 20b é mais barato). Conferir em groq.com/pricing"),
        ("Preço saída (US$ / 1M tokens)", 0.60,
         "Premissa: preço do gpt-oss-120b na Groq (teto; o 20b é mais barato). Conferir em groq.com/pricing"),
        ("Câmbio (R$ / US$)", 5.50, "Premissa: atualizar com a cotação do dia da coleta"),
    ]
    for r, (k, v, note) in enumerate(params, start=2):
        ws.cell(row=r, column=1, value=k).font = F_BOLD
        c = ws.cell(row=r, column=2, value=v)
        _style(c, "input")
        ws.cell(row=r, column=4, value=note).font = F_BASE

    ws["A8"] = "Custo dos concorrentes"
    ws["A8"].font = F_BOLD
    _header(ws, 9, ["Ferramenta", "Assinatura mensal (R$)", "Consultas / mês estimadas",
                    "Custo por consulta (R$)"])
    for r, tool in enumerate([t for t in ferramentas if not is_rag(t)], start=10):
        _style(ws.cell(row=r, column=1, value=tool), "key")
        _style(ws.cell(row=r, column=2), "input")
        _style(ws.cell(row=r, column=3), "input")
        c = ws.cell(row=r, column=4,
                    value=f'=IF(AND(ISNUMBER(B{r}),ISNUMBER(C{r}),C{r}>0),B{r}/C{r},"")')
        _style(c, "formula")
        c.number_format = "0.00"
    ws.cell(row=10 + len([t for t in ferramentas if not is_rag(t)]) + 1, column=1,
            value="O custo por consulta dos RAGs vem dos tokens medidos (aba Respostas).").font = F_BASE


def build_resumo(wb: Workbook, perguntas: list[dict], ferramentas: list[str], last: int) -> None:
    ws = wb.create_sheet("Resumo")
    ws["A1"] = "Resumo do benchmark"
    ws["A1"].font = F_TITLE

    heads = ["Ferramenta", "Avaliadas", "Nota média (0-10)", "Veredito", "Fundamentação",
             "Norma CE", "Acionável", "Sem alucinação", "% notas ≥ 8",
             "Latência média (s)", "Custo / consulta (R$)"]
    _header(ws, 3, heads)
    ws.column_dimensions["A"].width = 18
    for col in "BCDEFGHIJK":
        ws.column_dimensions[col].width = 14
    ws.row_dimensions[3].height = 30

    rng = lambda col: f"Respostas!${col}$2:${col}${last}"  # noqa: E731
    comp = [t for t in ferramentas if not is_rag(t)]
    for r, tool in enumerate(ferramentas, start=4):
        crit = {"D": "M", "E": "N", "F": "O", "G": "P", "H": "Q"}
        cells = {
            "A": tool,
            "B": f'=COUNTIFS({rng("B")},$A{r},{rng("R")},">=0")',
            "C": f'=IF(B{r}=0,"",AVERAGEIFS({rng("R")},{rng("B")},$A{r}))',
            "I": f'=IF(B{r}=0,"",COUNTIFS({rng("B")},$A{r},{rng("R")},">=8")/B{r})',
            "J": (f'=IF(COUNTIFS({rng("B")},$A{r},{rng("I")},">=0")=0,"",'
                  f'AVERAGEIFS({rng("I")},{rng("B")},$A{r}))'),
        }
        # Só linhas com nota completa, para bater com a "Nota média"
        for dst, src in crit.items():
            cells[dst] = (f'=IF(B{r}=0,"",'
                          f'AVERAGEIFS({rng(src)},{rng("B")},$A{r},{rng("R")},">=0"))')
        if is_rag(tool):
            cells["K"] = (f'=IF(COUNTIFS({rng("B")},$A{r},{rng("L")},">=0")=0,"",'
                          f'AVERAGEIFS({rng("L")},{rng("B")},$A{r}))')
        else:
            prow = 10 + comp.index(tool)
            cells["K"] = f'=IF(ISNUMBER(Parametros!$D${prow}),Parametros!$D${prow},"")'
        for col, v in cells.items():
            c = ws[f"{col}{r}"]
            c.value = v
            _style(c, "key" if col == "A" else "formula")
            if col in "CDEFGHJ" and col != "A":
                c.number_format = "0.00"
            if col == "I":
                c.number_format = "0%"
            if col == "K":
                c.number_format = "0.0000"
    last_tool_row = 3 + len(ferramentas)
    comp_cells = ",".join(f"C{4 + ferramentas.index(t)}" for t in comp)

    # Gate 1
    g = last_tool_row + 2
    ws.cell(row=g, column=1, value="Gate 1: sinal do benchmark").font = F_BOLD
    ws.cell(row=g + 1, column=1, value="Melhor concorrente").font = F_BASE
    best = ws.cell(row=g + 1, column=2,
                   value=f'=IF(COUNT({comp_cells})=0,"",MAX({comp_cells}))')
    _style(best, "formula")
    best.number_format = "0.00"
    ws.cell(row=g + 2, column=1, value="Sinal").font = F_BASE
    sinal = ws.cell(
        row=g + 2, column=2,
        value=(f'=IF(B{g + 1}="","Aguardando coleta",IF(B{g + 1}>=8,'
               '"Concorrentes fortes: só seguir se as entrevistas mostrarem dor mensurável",'
               f'IF(B{g + 1}<6,"Gap confirmado: seguir com H1 + H3",'
               '"Zona cinza: focar em H3 + H4 (fluxo + custo)")))'))
    _style(sinal, "formula")
    ws.merge_cells(start_row=g + 2, start_column=2, end_row=g + 2, end_column=8)
    ws.cell(row=g + 3, column=1,
            value="Regras em docs/fase3/03_plano_validacao.md. O Gate combina este sinal com as entrevistas.").font = F_BASE

    # Matriz categoria x ferramenta
    m = g + 5
    ws.cell(row=m, column=1, value="Nota média por categoria").font = F_BOLD
    _header(ws, m + 1, ["Categoria", *ferramentas])
    categorias = list(dict.fromkeys(p["categoria"] for p in perguntas))
    for i, cat in enumerate(categorias, start=m + 2):
        _style(ws.cell(row=i, column=1, value=cat), "key")
        for j, _ in enumerate(ferramentas, start=2):
            col = ws.cell(row=m + 1, column=j).column_letter
            c = ws.cell(
                row=i, column=j,
                value=(f'=IF(COUNTIFS({rng("B")},{col}${m + 1},{rng("C")},$A{i},{rng("R")},">=0")=0,"",'
                       f'AVERAGEIFS({rng("R")},{rng("B")},{col}${m + 1},{rng("C")},$A{i}))'))
            _style(c, "formula")
            c.number_format = "0.0"


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--force", action="store_true", help="sobrescreve a planilha existente")
    ap.add_argument("--out", type=Path, default=XLSX_PATH)
    args = ap.parse_args()

    if args.out.exists() and not args.force:
        print(f"ERRO: {args.out} já existe. Use --force para sobrescrever (apaga o preenchido).",
              file=sys.stderr)
        return 1

    data = json.loads(PERGUNTAS_JSON.read_text(encoding="utf-8"))
    perguntas, ferramentas = data["perguntas"], data["ferramentas"]
    ids = [p["id"] for p in perguntas]
    if len(ids) != len(set(ids)):
        print("ERRO: IDs duplicados em perguntas.json", file=sys.stderr)
        return 1
    if RAG_TOOL not in ferramentas:
        print(f"ERRO: '{RAG_TOOL}' precisa estar em ferramentas", file=sys.stderr)
        return 1

    wb = Workbook()
    build_leia_me(wb, len(perguntas), ferramentas)
    build_perguntas(wb, perguntas)
    last = build_respostas(wb, perguntas, ferramentas)
    build_parametros(wb, ferramentas)
    build_resumo(wb, perguntas, ferramentas, last)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print(f"OK: {args.out} ({len(perguntas)} perguntas × {len(ferramentas)} ferramentas = {last - 1} linhas)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
